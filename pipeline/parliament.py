"""Read-only Bundestag sources. No LLMs, inferred ballots, or credentials in snapshots."""
from __future__ import annotations

import io
import json
import os
import re
import time
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlencode, urljoin, urlsplit

import httpx
from bs4 import BeautifulSoup, NavigableString
from bs4.element import TemplateString
from openpyxl import load_workbook

from pipeline.models import GroupVote, MemberVote, Source
from pipeline.store import digest, write_json

DIP = 'https://search.dip.bundestag.de/api/v1'
DIP_HELP = 'https://content.dip.bundestag.de/content-api/v1/content/help-api'
BT = 'https://www.bundestag.de'
ROLL_CALL_INDEX = BT + '/ajax/filterlist/de/parlament/plenum/abstimmung/484422-484422'
BALLOT_INDEX = BT + '/ajax/filterlist/de/parlament/plenum/abstimmung/liste/462112-462112'
HOSTS = {'search.dip.bundestag.de', 'content.dip.bundestag.de',
         'www.bundestag.de', 'dserver.bundestag.de'}
CHOICES = ('yes', 'no', 'abstain', 'absent', 'invalid')
PARTIES = {'CDU/CSU': 'cdu-csu', 'SPD': 'spd', 'AfD': 'afd', 'FDP': 'fdp',
           'BÜNDNIS 90/DIE GRÜNEN': 'gruene', 'Bündnis 90/Die Grünen': 'gruene', 'BÜ90/GR': 'gruene',
           'B90/GRÜNE': 'gruene', 'B90/Grüne': 'gruene',
           'DIE LINKE': 'linke', 'Die Linke': 'linke', 'BSW': 'bsw', 'SSW': 'ssw'}


class SourceError(RuntimeError):
    """Safe, bounded diagnostics: never include authorization headers or response bodies."""


def normalized(text: str) -> str:
    return re.sub(r'\W+', '', unicodedata.normalize('NFKC', text).casefold().replace('\xad', ''))


def document_numbers(text: str) -> set[str]:
    return {f'{term}/{number}' for term, number in re.findall(r'\b(\d{1,2})/\s*(\d{1,6})\b', text)}


def official_url(url: str) -> str:
    parsed = urlsplit(url)
    if (parsed.scheme != 'https' or parsed.hostname not in HOSTS or parsed.username
            or parsed.password or parsed.port not in (None, 443)
            or re.search(r'(?i)(apikey|token|authorization)', parsed.query)):
        raise SourceError('Untrusted parliamentary source URL')
    return url.split('#')[0]


@dataclass
class Download:
    url: str
    content: bytes
    retrieved_at: date

    def source(self, title: str) -> Source:
        supplementary = urlsplit(self.url).hostname == 'www.abgeordnetenwatch.de'
        return Source(url=self.url, title=title,
                      publisher='abgeordnetenwatch.de' if supplementary else 'Deutscher Bundestag',
                      retrieved_at=self.retrieved_at, sha256=digest(self.content),
                      license_note='CC0 1.0; ergänzender, nicht amtlicher Abgleich.' if supplementary else
                                   'Amtlicher parlamentarischer Nachweis; Quelle und Kontext beibehalten.')

    def json(self):
        try:
            return json.loads(self.content)
        except (ValueError, UnicodeError) as error:
            raise SourceError('Invalid parliamentary JSON response') from error


class BundestagClient:
    checked_url = staticmethod(official_url)
    request_interval = 0.3

    def __init__(self, cache: Path, *, refresh=False, client=None, sleep=time.sleep, clock=time.monotonic):
        self.cache, self.refresh = cache, refresh
        self.http = client or httpx.Client(timeout=60, follow_redirects=False,
            headers={'User-Agent': 'Politrace/1.0 (https://github.com/StromFLIX/politrace)'})
        self.sleep, self.clock = sleep, clock
        self._key = None
        self._key_error = None
        self._last_request = None
        self._refreshed = set()
        self._failed = {}

    def close(self):
        self.http.close()

    def api_key(self):
        if self._key_error:
            raise self._key_error
        if not self._key:
            self._key = os.environ.get('DIP_API_KEY', '').strip()
            if not self._key:
                # This is the documented shared public key, not the portal's private config.
                # Keep it in memory only; do not cache the help response or use query auth.
                try:
                    response = self._request(DIP_HELP)
                except SourceError:
                    self._key_error = SourceError('DIP public API documentation unavailable; no access key could be obtained')
                    raise self._key_error from None
                match = re.search(rb'\b[A-Za-z0-9]{7}\.[A-Za-z0-9]{34}\b', response)
                if not match:
                    raise SourceError('DIP public access key unavailable; configure DIP_API_KEY in the job environment')
                self._key = match[0].decode('ascii')
        return self._key

    def _request(self, url: str) -> bytes:
        url = self.checked_url(url)
        for attempt in range(3):
            headers = {'Authorization': 'ApiKey ' + self.api_key()} if url.startswith(DIP + '/') else {}
            if self._last_request is not None:
                self.sleep(max(0, self.request_interval - (self.clock() - self._last_request)))
            self._last_request = self.clock()
            retry_after = 2 ** attempt
            try:
                with self.http.stream('GET', url, headers=headers) as response:
                    status = response.status_code
                    if status == 200:
                        content = bytearray()
                        for chunk in response.iter_bytes():
                            content.extend(chunk)
                            if len(content) > 40_000_000:
                                raise SourceError('Parliamentary source exceeds the download limit')
                        return bytes(content)
                    if status == 303 and '/.enodia/challenge' in response.headers.get('Location', ''):
                        raise SourceError('Official source requires an interactive browser challenge (HTTP 303)')
                    retry = status == 429 or status >= 500
                    if status == 429:
                        wait = response.headers.get('Retry-After', '')
                        retry_after = min(120, max(2, int(wait))) if wait.isdigit() else 60
            except httpx.HTTPError:
                status, retry = 'network', True
            if not retry or attempt == 2:
                raise SourceError(f'Parliamentary source request failed (HTTP {status})') from None
            self.sleep(retry_after)
        raise SourceError('Parliamentary source retry limit reached')

    def get(self, url: str, params=None) -> Download:
        url = self.checked_url(url)
        if params:
            url += ('&' if '?' in url else '?') + urlencode(sorted(params.items()), doseq=True)
        url = self.checked_url(url)
        key = digest(url)
        blob, meta = self.cache / f'{key}.bin', self.cache / f'{key}.json'
        # Bounded daily snapshots for settled parliamentary evidence; fast-expiring
        # discovery indices and negative API lookups. Retrieval dates are preserved.
        ttl = 3600 if '/ajax/' in url or url.endswith('/services/opendata') else 86400
        if (not self.refresh or url in self._refreshed) and blob.exists() and meta.exists():
            metadata = json.loads(meta.read_text())
            content = blob.read_bytes()
            if '/api/' in url:
                try:
                    body = json.loads(content)
                    if body.get('numFound') == 0 or body.get('meta', {}).get('result', {}).get('total') == 0:
                        ttl = 3600
                except (ValueError, AttributeError):
                    ttl = 0
            if (time.time() - meta.stat().st_mtime < ttl and metadata['url'] == url
                    and metadata['sha256'] == digest(content)):
                return Download(url, content, date.fromisoformat(metadata['retrieved_at']))
        if url in self._failed:
            raise self._failed[url]
        try:
            content = self._request(url)
        except SourceError as error:
            self._failed[url] = error  # Avoid hammering the same failed source within a run.
            raise
        downloaded = Download(url, content, date.today())
        self.cache.mkdir(parents=True, exist_ok=True)
        temporary = blob.with_suffix('.tmp')
        temporary.write_bytes(content)
        temporary.replace(blob)
        write_json(meta, {'url': url, 'sha256': digest(content), 'retrieved_at': str(downloaded.retrieved_at)})
        self._refreshed.add(url)
        return downloaded

    def listing(self, resource: str, params: dict) -> list[dict]:
        documents, seen, cursors = [], set(), set()
        cursor, expected = None, None
        for _ in range(100):
            query = {**params, **({'cursor': cursor} if cursor else {})}
            body = self.get(f'{DIP}/{resource}', query).json()
            if not isinstance(body.get('numFound'), int) or not isinstance(body.get('documents'), list):
                raise SourceError('Invalid DIP list envelope')
            expected = body['numFound'] if expected is None else expected
            if body['numFound'] != expected:
                raise SourceError('DIP inventory changed during pagination; retry the scan')
            for item in body['documents']:
                if not str(item.get('id', '')).isdigit() or item['id'] in seen:
                    raise SourceError('Duplicate or invalid DIP identity during pagination')
                seen.add(item['id'])
                documents.append(item)
            if len(documents) == expected:
                return documents
            next_cursor = body.get('cursor')
            if not body['documents'] or not next_cursor or next_cursor in cursors or len(documents) > expected:
                raise SourceError('Incomplete DIP pagination')
            cursors.add(next_cursor)
            cursor = next_cursor
        raise SourceError('DIP pagination limit reached')


def parse_roll_calls(content: bytes) -> tuple[list[dict], int, int]:
    soup = BeautifulSoup(content, 'html.parser')
    meta = soup.select_one('.meta-slider[data-hits][data-nextoffset]')
    if meta is None:
        raise SourceError('Bundestag roll-call index format changed')
    entries = []
    for slide in soup.select('.bt-slide'):
        chart, stamp = slide.select_one('canvas[id^="canvas-na-"]'), slide.select_one('.bt-date')
        heading = slide.select_one('.bt-teaser-text h3')
        motion = slide.select_one('.bt-teaser-haupttext')
        if chart is None or stamp is None or heading is None or motion is None:
            raise SourceError('Incomplete official roll-call index entry')
        leading = heading.select_one('.bt-dachzeile')
        if leading:
            leading.decompose()
        counts = [int(n) for n in chart['data-chart-values'].split(',')]
        if len(counts) != 4 or any(n < 0 for n in counts):
            raise SourceError('Unexpected official roll-call total columns')
        entries.append({'id': chart['id'].removeprefix('canvas-na-'),
            'date': datetime.strptime(stamp.get_text(strip=True), '%d.%m.%Y').date(),
            'title': heading.get_text(' ', strip=True), 'motion': motion.get_text(' ', strip=True),
            'totals': dict(zip(CHOICES, [*counts, 0]))})
    return entries, int(meta['data-nextoffset']), int(meta['data-hits'])


def roll_call_inventory(client: BundestagClient, since: date) -> list[dict]:
    result, seen, offset = [], set(), 0
    for _ in range(150):
        rows, next_offset, total = parse_roll_calls(client.get(ROLL_CALL_INDEX, {'offset': offset}).content)
        if not rows or next_offset <= offset:
            raise SourceError('Incomplete Bundestag roll-call pagination')
        for row in rows:
            if row['id'] in seen:
                raise SourceError('Repeated Bundestag roll-call index page')
            seen.add(row['id'])
            if row['date'] >= since:
                result.append(row)
        if min(row['date'] for row in rows) < since or next_offset >= total:
            return result
        offset = next_offset
    raise SourceError('Bundestag roll-call pagination limit reached')


def parse_ballot_links(content: bytes) -> list[dict]:
    soup = BeautifulSoup(content, 'html.parser')
    table = soup.select_one('template[data-js-document-results="table"]')
    if table is None:
        raise SourceError('Bundestag ballot download index format changed')
    result = []
    for row in table.select('tr'):
        pdf = row.select_one('a[href$=".pdf"]')
        sheet = row.select_one('a[href$=".xlsx"]')
        if pdf is None or sheet is None:
            raise SourceError('Missing official ballot download metadata')
        match = re.search(r'/(\d{8})[_-](?:(\d+)(?:[_-][^/]*)?|xls)\.xlsx$', sheet['href'])
        if not match:
            raise SourceError('Unrecognized Bundestag ballot date')
        title = (pdf.find('span') or pdf).get_text(' ', strip=True, types=(NavigableString, TemplateString))
        if not title:
            raise SourceError('Missing official ballot download title')
        title = re.sub(r'^\d{2}\.\d{2}\.\d{4}:\s*', '', title)
        result.append({'date': datetime.strptime(match[1], '%Y%m%d').date(), 'title': title,
                       'url': official_url(urljoin(BT, sheet['href']))})
    return result


def ballot_inventory(client: BundestagClient, since: date) -> list[dict]:
    result, seen = [], set()
    for offset in range(0, 3000, 30):
        downloaded = client.get(BALLOT_INDEX, {'limit': 30, 'offset': offset})
        rows = parse_ballot_links(downloaded.content)
        for row in rows:
            row['index_source'] = downloaded.source('Bundestag: amtliche Abstimmungslisten')
            if row['url'] in seen:
                raise SourceError('Repeated Bundestag ballot download page')
            seen.add(row['url'])
            if row['date'] >= since:
                result.append(row)
        if len(rows) < 30 or min(row['date'] for row in rows) < since:
            return result
    raise SourceError('Bundestag ballot download pagination limit reached')


def group_key(group: str) -> str:
    key = normalized(group)
    if key in {'fraktionslos', 'fraktionslose'}:
        return 'independent'
    return next((party for label, party in PARTIES.items() if normalized(label) == key), key)


def group_tallies(groups: list[GroupVote]) -> dict:
    result = {}
    for group in groups:
        key = group_key(group.group)
        if key in result:
            raise SourceError('Duplicate normalized voting group')
        result[key] = {choice: getattr(group, choice) for choice in CHOICES}
    return result


def roll_call_detail(download: Download, entry: dict) -> dict:
    soup = BeautifulSoup(download.content, 'html.parser')
    heading = soup.select_one('h1.bt-artikel__title')
    motion = heading.find_next_sibling('p') if heading else None
    if heading is None or motion is None:
        raise SourceError('Official roll-call detail has no complete motion')
    stamp = heading.find_previous_sibling('span', class_='bt-date')
    match = re.fullmatch(r'(\d{1,2})\.\s+(\w+)\s+(\d{4})', stamp.get_text(' ', strip=True) if stamp else '')
    months = ['Januar', 'Februar', 'März', 'April', 'Mai', 'Juni', 'Juli', 'August',
              'September', 'Oktober', 'November', 'Dezember']
    if not match or match[2] not in months or date(int(match[3]), months.index(match[2]) + 1, int(match[1])) != entry['date']:
        raise SourceError('Official roll-call detail disagrees with the dated index')
    overall = soup.select('canvas[data-chart-type="bar"][data-chart-values]')
    if len(overall) != 1 or overall[0]['data-chart-values'] != ','.join(str(entry['totals'][c]) for c in CHOICES[:4]):
        raise SourceError('Official structured result disagrees with the roll-call index')
    # Some detail pages additionally embed JSON results in the debate report. Check
    # them when present; the primary structured faction chart is always required.
    results = []
    for script in soup.select('script[type="application/json"]'):
        text = script.get_text().strip().removeprefix('<![CDATA[').removesuffix(']]>')
        try:
            data = json.loads(text)
        except ValueError:
            continue
        if isinstance(data, dict) and data.get('href', '').endswith('?id=' + entry['id']):
            results.append(data)
    if any(result.get('date') != str(entry['date']) or
                          {**result.get('votes', {}), 'invalid': 0} != entry['totals'] for result in results):
        raise SourceError('Official roll-call JSON result disagrees with the dated index')
    groups = []
    for element in soup.select('.bt-teaser-chart-solo[data-value]'):
        chart = element.select_one('canvas[data-chart-values]')
        if chart is None:
            raise SourceError('Missing official faction tally')
        counts = [int(value) for value in chart['data-chart-values'].split(',')]
        if len(counts) != 4 or any(n < 0 for n in counts):
            raise SourceError('Unexpected official faction result columns')
        groups.append(GroupVote(group=element['data-value'], **dict(zip(CHOICES, [*counts, 0]))))
    if not groups or {c: sum(getattr(g, c) for g in groups) for c in CHOICES} != entry['totals']:
        raise SourceError('Official faction subtotals do not reconcile with the overall result')
    return {**entry, 'title': heading.get_text(' ', strip=True), 'motion': motion.get_text(' ', strip=True),
            'group_totals': group_tallies(groups), 'download': download}


def parse_ballots(content: bytes, *, term: int, sitting: int, expected: dict | None = None,
                  ballot: int | None = None):
    """One selected option per row; reconcile *all* ballots, not just yes/no majorities."""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 20_000_000:
            raise SourceError('Oversized ballot workbook')
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) != 1:
            raise SourceError('Unexpected multiple-sheet ballot workbook')
        sheet = workbook.active
        if sheet.max_row > 2000 or sheet.max_column > 40:
            raise SourceError('Unexpected ballot sheet dimensions')
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() for value in next(rows)]
        required = ['Wahlperiode', 'Sitzungnr', 'Abstimmnr', 'Fraktion/Gruppe', 'Name', 'Vorname',
                    'ja', 'nein', 'Enthaltung', 'nichtabgegeben', 'ungültig']
        if len(headers) != len(set(headers)) or not set(required) <= set(headers):
            raise SourceError('Unknown ballot workbook columns')
        members, counts, seen, ballots = [], {}, set(), set()
        for row_number, values in enumerate(rows, 2):
            if not any(value is not None for value in values):
                continue
            if any(isinstance(value, str) and value.startswith('=') for value in values):
                raise SourceError('Formula-based ballot row')
            row = dict(zip(headers, values))
            if row['Wahlperiode'] != term or row['Sitzungnr'] != sitting:
                raise SourceError('Ballot workbook does not belong to the DIP sitting')
            number = row['Abstimmnr']
            if type(number) not in (int, float) or not 1 <= number <= 1000 or int(number) != number:
                raise SourceError('Missing or invalid workbook ballot number')
            if ballot is not None and number != ballot:
                raise SourceError('Ballot workbook does not belong to the identified decision')
            ballots.add(int(number))
            group = str(row['Fraktion/Gruppe'] or '').strip()
            name = str(row.get('Bezeichnung') or f"{row['Vorname'] or ''} {row['Name'] or ''}").strip()
            if not group or not row['Name'] or (group, name) in seen:
                raise SourceError('Missing or duplicate member ballot identity')
            seen.add((group, name))
            flags = [row[key] if row[key] is not None else 0
                     for key in ('ja', 'nein', 'Enthaltung', 'nichtabgegeben', 'ungültig')]
            if any(type(value) not in (int, float) or value not in (0, 1) for value in flags) or sum(flags) != 1:
                raise SourceError('Invalid, formula-based, or ambiguous ballot row')
            choice = CHOICES[flags.index(1)]
            members.append(MemberVote(name=name, group=group, vote=choice, source_row=row_number))
            counts.setdefault(group, dict.fromkeys(CHOICES, 0))[choice] += 1
        if len(ballots) != 1 or not members:
            raise SourceError('Empty or mixed-decision ballot workbook')
        totals = {choice: sum(group[choice] for group in counts.values()) for choice in CHOICES}
        if expected is not None and totals != expected:
            raise SourceError('Individual ballots disagree with the official roll-call totals')
        groups = [GroupVote(group=group, party_id=PARTIES.get(group), **count)
                  for group, count in counts.items()]
        return members, groups, next(iter(ballots))
    finally:
        workbook.close()
